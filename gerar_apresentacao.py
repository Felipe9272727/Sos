# -*- coding: utf-8 -*-
"""Gera uma apresentacao .pptx sobre Ultimate Frisbee."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- Paleta de cores ----
AZUL_ESCURO = RGBColor(0x0F, 0x2C, 0x4A)   # fundo principal
AZUL        = RGBColor(0x1B, 0x6C, 0xA8)   # detalhes
LARANJA     = RGBColor(0xFF, 0x8C, 0x1A)   # destaque (cor do disco)
AMARELO     = RGBColor(0xFF, 0xC4, 0x3D)
BRANCO      = RGBColor(0xFF, 0xFF, 0xFF)
CINZA_CLARO = RGBColor(0xEC, 0xF2, 0xF7)
TEXTO_ESC   = RGBColor(0x16, 0x2A, 0x3A)

prs = Presentation()
prs.slide_width = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
LARG = prs.slide_width
ALT = prs.slide_height
BLANK = prs.slide_layouts[6]

def fundo(slide, cor):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, LARG, ALT)
    s.fill.solid(); s.fill.fore_color.rgb = cor
    s.line.fill.background()
    s.shadow.inherit = False
    # manda pro fundo
    slide.shapes._spTree.remove(s._element)
    slide.shapes._spTree.insert(2, s._element)
    return s

def faixa(slide, x, y, w, h, cor):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = cor
    s.line.fill.background()
    s.shadow.inherit = False
    return s

def circulo(slide, x, y, d, cor):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    s.fill.solid(); s.fill.fore_color.rgb = cor
    s.line.fill.background()
    s.shadow.inherit = False
    return s

def caixa_texto(slide, x, y, w, h, texto, tam=20, cor=TEXTO_ESC, bold=False,
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fonte="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run(); r.text = texto
    r.font.size = Pt(tam); r.font.bold = bold
    r.font.color.rgb = cor; r.font.name = fonte
    return tb

def lista(slide, x, y, w, h, itens, tam=18, cor=TEXTO_ESC, cor_marcador=LARANJA,
          espaco=8):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(itens):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(espaco)
        r = p.add_run(); r.text = "●  "
        r.font.size = Pt(tam); r.font.color.rgb = cor_marcador; r.font.bold = True
        r2 = p.add_run(); r2.text = item
        r2.font.size = Pt(tam); r2.font.color.rgb = cor; r2.font.name = "Calibri"
    return tb

def titulo_secao(slide, texto, n):
    # faixa lateral + numero
    faixa(slide, 0, 0, Inches(0.35), ALT, LARANJA)
    caixa_texto(slide, Inches(0.7), Inches(0.35), Inches(11.5), Inches(1.0),
                texto, tam=34, cor=AZUL_ESCURO, bold=True)
    faixa(slide, Inches(0.75), Inches(1.35), Inches(3.2), Pt(4), LARANJA)
    # numero do slide canto
    caixa_texto(slide, LARG-Inches(1.2), ALT-Inches(0.7), Inches(0.9), Inches(0.5),
                str(n), tam=14, cor=AZUL, bold=True, align=PP_ALIGN.RIGHT)

# ============ SLIDE 1 - CAPA ============
s = prs.slides.add_slide(BLANK)
fundo(s, AZUL_ESCURO)
# circulos decorativos (discos)
circulo(s, Inches(-1.5), Inches(-1.5), Inches(4), AZUL)
circulo(s, LARG-Inches(2.8), ALT-Inches(2.8), Inches(4.5), LARANJA)
circulo(s, LARG-Inches(2.0), ALT-Inches(2.0), Inches(2.9), AZUL_ESCURO)
# disco central estilizado
circulo(s, Inches(5.5), Inches(0.6), Inches(2.3), LARANJA)
circulo(s, Inches(5.85), Inches(0.95), Inches(1.6), AMARELO)
caixa_texto(s, Inches(5.5), Inches(1.25), Inches(2.3), Inches(1.0),
            "🥏", tam=54, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
caixa_texto(s, Inches(1), Inches(3.1), Inches(11.3), Inches(1.2),
            "ULTIMATE FRISBEE", tam=54, cor=BRANCO, bold=True, align=PP_ALIGN.CENTER)
caixa_texto(s, Inches(1), Inches(4.3), Inches(11.3), Inches(0.7),
            "Origem, regras e curiosidades do esporte", tam=22, cor=AMARELO,
            align=PP_ALIGN.CENTER)
caixa_texto(s, Inches(1), Inches(6.1), Inches(11.3), Inches(0.8),
            "Nicolas  •  Felipe  •  André  •  Lucas  •  Ruan", tam=20, cor=BRANCO,
            align=PP_ALIGN.CENTER, bold=True)

# ============ SLIDE 2 - O QUE É / ORIGEM ============
s = prs.slides.add_slide(BLANK)
fundo(s, CINZA_CLARO)
titulo_secao(s, "O que é e como surgiu", 2)
caixa_texto(s, Inches(0.75), Inches(1.7), Inches(11.8), Inches(1.4),
            "O Ultimate Frisbee é um esporte coletivo jogado com um disco (frisbee). "
            "O objetivo é levar o disco até a área de pontuação do adversário por meio de passes. "
            "Começou como brincadeira universitária e virou um esporte competitivo praticado no mundo todo.",
            tam=19, cor=TEXTO_ESC)
# bloco origem do nome
faixa(s, Inches(0.75), Inches(3.5), Inches(11.8), Inches(3.2), BRANCO)
faixa(s, Inches(0.75), Inches(3.5), Inches(0.18), Inches(3.2), LARANJA)
caixa_texto(s, Inches(1.1), Inches(3.65), Inches(11.2), Inches(0.6),
            "📜 De onde vem o nome \"Frisbee\"?", tam=22, cor=AZUL_ESCURO, bold=True)
caixa_texto(s, Inches(1.1), Inches(4.35), Inches(11.2), Inches(2.2),
            "O frisbee se inspirou em formas antigas de arremessar discos, mas ficou famoso "
            "nos EUA nos anos 1940. O nome veio da empresa de tortas \"Frisbie Pie Company\": "
            "os estudantes jogavam as formas de torta vazias gritando \"Frisbie!\". "
            "Depois, a empresa Wham-O criou o disco de plástico moderno e mudou a escrita para \"Frisbee\".",
            tam=18, cor=TEXTO_ESC)

# ============ SLIDE 3 - REGRAS BÁSICAS ============
s = prs.slides.add_slide(BLANK)
fundo(s, CINZA_CLARO)
titulo_secao(s, "Regras básicas", 3)
lista(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(5),
      [
        "O objetivo é chegar à área de pontuação do adversário recebendo o disco.",
        "Quem está com o frisbee NÃO pode correr — só pode passar.",
        "Cada equipe joga normalmente com 7 jogadores em campo.",
        "O disco passa para o outro time se cair no chão ou for interceptado.",
        "Pontua-se quando um jogador recebe o disco dentro da área de gol adversária.",
        "Não é permitido contato físico forte entre os jogadores.",
      ], tam=21, espaco=14)

# ============ SLIDE 4 - DINÂMICA DO JOGO ============
s = prs.slides.add_slide(BLANK)
fundo(s, CINZA_CLARO)
titulo_secao(s, "Como o jogo acontece", 4)
lista(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(5),
      [
        "O jogo começa com um lançamento chamado \"pull\", parecido com a saída do futebol.",
        "Quem está com o disco tem cerca de 10 segundos para passar.",
        "Se o disco sair do campo, a posse vai para o outro time.",
        "Após cada ponto, os times trocam de lado no campo.",
        "As substituições normalmente acontecem após um ponto ser marcado.",
        "Se dois jogadores pegam o disco ao mesmo tempo, a posse fica com o ataque.",
        "Pode ser jogado em grama, areia ou quadras adaptadas.",
        "Em competições oficiais, vence quem chega primeiro à pontuação definida (geralmente 15 pontos).",
      ], tam=18, espaco=10)

# ============ SLIDE 5 - ESPÍRITO DO JOGO ============
s = prs.slides.add_slide(BLANK)
fundo(s, AZUL_ESCURO)
circulo(s, LARG-Inches(3.0), Inches(-1.2), Inches(4), AZUL)
circulo(s, Inches(-1.2), ALT-Inches(2.8), Inches(4), LARANJA)
caixa_texto(s, Inches(1), Inches(1.3), Inches(11.3), Inches(1.0),
            "🤝 O Espírito do Jogo", tam=40, cor=AMARELO, bold=True, align=PP_ALIGN.CENTER)
caixa_texto(s, Inches(1.5), Inches(3.0), Inches(10.3), Inches(2.5),
            "O Ultimate Frisbee é conhecido por não ter árbitros na maioria das partidas: "
            "os próprios jogadores resolvem as faltas. Por isso, o esporte valoriza o "
            "\"Espírito do Jogo\", que incentiva respeito, honestidade e fair play entre todos.",
            tam=24, cor=BRANCO, align=PP_ALIGN.CENTER)

# ============ SLIDE 6 - CURIOSIDADES ============
s = prs.slides.add_slide(BLANK)
fundo(s, CINZA_CLARO)
titulo_secao(s, "Curiosidades", 6)
curiosidades = [
    ("⚡", "O frisbee pode ultrapassar 100 km/h dependendo do lançamento."),
    ("🏆", "Existe campeonato mundial de frisbee."),
    ("🐶", "Cachorros também competem em provas de frisbee!"),
    ("🧠", "O esporte ajuda nos reflexos, na coordenação e no trabalho em equipe."),
    ("🎯", "Há várias modalidades: freestyle, disc golf e ultimate frisbee."),
]
y = Inches(1.95)
for emoji, txt in curiosidades:
    card = faixa(s, Inches(0.9), y, Inches(11.5), Inches(0.85), BRANCO)
    caixa_texto(s, Inches(1.05), y, Inches(0.9), Inches(0.85), emoji, tam=26,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    caixa_texto(s, Inches(2.0), y, Inches(10.2), Inches(0.85), txt, tam=18,
                cor=TEXTO_ESC, anchor=MSO_ANCHOR.MIDDLE)
    y = y + Inches(1.0)

# ============ SLIDE 7 - ENCERRAMENTO ============
s = prs.slides.add_slide(BLANK)
fundo(s, AZUL_ESCURO)
circulo(s, Inches(5.4), Inches(1.6), Inches(2.5), LARANJA)
circulo(s, Inches(5.78), Inches(1.98), Inches(1.74), AMARELO)
caixa_texto(s, Inches(5.4), Inches(2.25), Inches(2.5), Inches(1.2),
            "🥏", tam=60, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
caixa_texto(s, Inches(1), Inches(4.4), Inches(11.3), Inches(1.0),
            "Obrigado pela atenção!", tam=44, cor=BRANCO, bold=True, align=PP_ALIGN.CENTER)
caixa_texto(s, Inches(1), Inches(5.7), Inches(11.3), Inches(0.8),
            "Nicolas  •  Felipe  •  André  •  Lucas  •  Ruan", tam=20, cor=AMARELO,
            align=PP_ALIGN.CENTER, bold=True)

prs.save("Ultimate_Frisbee.pptx")
print("OK - arquivo gerado com", len(prs.slides._sldIdLst), "slides")
